from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEAD = [
    "№",
    "ФИО",
    "Должность",
    "Наименование KPI",
    "Единица измерения",
    "План",
    "Факт",
    "Исполнение показателя (%)",
    "Вес (%)",
    "Исполнение личного KPI (%)",
    "Подтверждающий материал",
    "Примечание",
]


def build_xlsx(state: dict) -> bytes:
    r = state["result"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Лист исполнения"
    ws.append(HEAD)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.fill = PatternFill("solid", fgColor="DDE6F3")
    for i, b in enumerate(r["blocks"], 1):
        unit = "%" if b["kind"] == "metric" else "шт."
        if b["kind"] == "count":
            ev = "; ".join(f"[{m['indicator_num']}] {m['quote']} ({m['source']})" for m in b["matches"] if m["counted"])
            note = "; ".join(
                f"НЕ ЗАСЧИТАНО [{m['indicator_num']}]: {m['missing_conditions']}" for m in b["matches"] if not m["counted"]
            )
        else:
            ev = "; ".join(f"[{m['indicator_num']}] {m['quote']} ({m['source']})" for m in b["metrics"])
            note = " ".join(b["notes"])
        ws.append(
            [
                i,
                r["employee"],
                r["position"]["name"],
                b["title"],
                unit,
                b["target"]["value"],
                b["achieved"],
                b["fulfillment"],
                b["weight"],
                b["contribution"],
                ev,
                note,
            ]
        )
    ws.append(
        [
            "",
            "",
            "",
            "Итого %",
            "",
            "",
            "",
            "",
            r.get("weights_total"),
            r["kpi_total_pct"],
            "",
            "Предварительная оценка; утверждается Комиссией по оценке KPI",
        ]
    )
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="FFF2CC")
    for col in range(1, len(HEAD) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 60 if col in (4, 11, 12) else 16
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
