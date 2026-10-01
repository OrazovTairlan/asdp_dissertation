from __future__ import annotations

import contextlib
import re
from dataclasses import dataclass, field
from pathlib import Path

NUM_ITEM = re.compile(r"^\s*\d{1,2}[.)]\s+\S")
BLOCK_LINE = re.compile(r"^\s*Блок\s*\d+", re.I)
MODE_LINE = re.compile(r"^\s*(Один из следующих|Среднее значение)", re.I)


@dataclass
class TableData:
    rows: list[list[str]]
    source: str
    page: int | None
    context: str = ""
    index: int = 0

    def to_markdown(self) -> str:
        if not self.rows:
            return ""
        width = max(len(r) for r in self.rows)
        rows = [[c.replace("\n", "<br>").replace("|", "\\|") for c in r] + [""] * (width - len(r)) for r in self.rows]
        head, body = rows[0], rows[1:]
        md = ["| " + " | ".join(head) + " |", "|" + "---|" * width]
        md += ["| " + " | ".join(r) + " |" for r in body]
        return "\n".join(md)


@dataclass
class ParsedDoc:
    filename: str
    text: str
    tables: list[TableData] = field(default_factory=list)
    needs_ocr_pages: list[int] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    def full_markdown(self) -> str:
        parts = [self.text.strip()]
        for t in self.tables:
            where = f"стр. {t.page}" if t.page else f"таблица {t.index + 1}"
            parts.append(f"\n[Таблица, {where}]\n" + t.to_markdown())
        return "\n\n".join(p for p in parts if p)


def join_wrapped(text: str | None) -> str:
    if not text:
        return ""
    out: list[str] = []
    for raw in text.replace("\r", "").split("\n"):
        line = raw.strip()
        if not line:
            continue
        if not out:
            out.append(line)
            continue
        prev = out[-1]
        starts_new = bool(NUM_ITEM.match(line) or BLOCK_LINE.match(line) or MODE_LINE.match(line))
        prev_is_head = bool(MODE_LINE.match(prev) and prev.endswith(":"))
        if starts_new or prev_is_head:
            out.append(line)
        elif prev.endswith("-") and line[:1].islower():
            out[-1] = prev + line
        else:
            out[-1] = prev + " " + line
    return "\n".join(out)


def _clean_cell(v) -> str:
    return join_wrapped(v if isinstance(v, str) else ("" if v is None else str(v)))


def _parse_pdf(path: Path) -> ParsedDoc:
    import pdfplumber

    doc = ParsedDoc(filename=path.name, text="")
    texts: list[str] = []
    tindex = 0
    with pdfplumber.open(path) as pdf:
        for pno, page in enumerate(pdf.pages, 1):
            found = _drop_nested(sorted(page.find_tables(), key=lambda t: t.bbox[1]))
            rest = page
            prev_bottom = 0.0
            for t in found:
                _, top, _, bottom = t.bbox
                ctx = ""
                if top - prev_bottom > 3:
                    try:
                        ctx = page.crop((0, max(prev_bottom, 0), page.width, min(top, page.height))).extract_text() or ""
                    except Exception:
                        ctx = ""
                rows = [[_clean_cell(c) for c in r] for r in t.extract()]
                rows = [r for r in rows if any(c for c in r)]
                if rows:
                    doc.tables.append(TableData(rows=rows, source=path.name, page=pno, context=_context_line(ctx), index=tindex))
                    tindex += 1
                prev_bottom = bottom
                with contextlib.suppress(Exception):
                    rest = rest.outside_bbox(t.bbox)
            page_text = (rest.extract_text() or "").strip()
            if not page_text and not found:
                doc.needs_ocr_pages.append(pno)
            texts.append(page_text)
    doc.text = "\n\n".join(t for t in texts if t)
    if doc.needs_ocr_pages:
        doc.warnings.append(f"Страницы без текстового слоя (скан): {doc.needs_ocr_pages} — нужен OCR через vision-модель.")
    return doc


def _drop_nested(tables: list) -> list:

    def inside(a, b) -> bool:
        return a[0] >= b[0] - 2 and a[1] >= b[1] - 2 and a[2] <= b[2] + 2 and a[3] <= b[3] + 2

    return [t for t in tables if not any(o is not t and inside(t.bbox, o.bbox) for o in tables)]


def _context_line(text: str) -> str:
    lines = [ln.strip() for ln in text.split("\n") if ln.strip() and not ln.strip().isdigit()]
    if not lines:
        return ""
    starts = [i for i, ln in enumerate(lines) if re.match(r"(?i)^для\s", ln)]
    if starts:
        return " ".join(lines[starts[-1] :])
    return lines[-1]


def _parse_docx(path: Path) -> ParsedDoc:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    d = docx.Document(str(path))
    doc = ParsedDoc(filename=path.name, text="")
    paras: list[str] = []
    last_para = ""
    tindex = 0
    for child in d.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            t = Paragraph(child, d).text.strip()
            if t:
                paras.append(t)
                last_para = t
        elif tag == "tbl":
            tbl = Table(child, d)
            rows: list[list[str]] = []
            for r in tbl.rows:
                cells, seen = [], set()
                for c in r.cells:
                    key = id(c._tc)
                    cells.append("" if key in seen else _clean_cell(c.text))
                    seen.add(key)
                rows.append(cells)
            rows = [r for r in rows if any(c for c in r)]
            if rows:
                doc.tables.append(TableData(rows=rows, source=path.name, page=None, context=last_para, index=tindex))
                tindex += 1
    doc.text = "\n".join(paras)
    return doc


def _parse_xlsx(path: Path) -> ParsedDoc:
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True)
    doc = ParsedDoc(filename=path.name, text="")
    for i, ws in enumerate(wb.worksheets):
        grid = [[("" if c.value is None else str(c.value)) for c in row] for row in ws.iter_rows()]
        for rng in ws.merged_cells.ranges:
            v = grid[rng.min_row - 1][rng.min_col - 1]
            for r in range(rng.min_row - 1, rng.max_row):
                for c in range(rng.min_col - 1, rng.max_col):
                    if (r, c) != (rng.min_row - 1, rng.min_col - 1):
                        grid[r][c] = ""
            grid[rng.min_row - 1][rng.min_col - 1] = v
        rows = [[_clean_cell(c) for c in r] for r in grid]
        rows = [r for r in rows if any(c for c in r)]
        if rows:
            doc.tables.append(TableData(rows=rows, source=path.name, page=None, context=ws.title, index=i))
    return doc


def parse_file(path: Path) -> ParsedDoc:
    ext = path.suffix.lower()
    if ext == ".pdf":
        return _parse_pdf(path)
    if ext == ".docx":
        return _parse_docx(path)
    if ext in (".xlsx", ".xlsm"):
        return _parse_xlsx(path)
    if ext in (".txt", ".md", ".csv"):
        return ParsedDoc(filename=path.name, text=path.read_text(encoding="utf-8", errors="replace"))
    raise ValueError(f"Неподдерживаемый формат: {ext}")


def render_pdf_pages(path: Path, pages: list[int] | None = None, scale: float = 2.0) -> list[tuple[int, bytes]]:
    import io

    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(str(path))
    out = []
    for i in range(len(pdf)):
        if pages and (i + 1) not in pages:
            continue
        img = pdf[i].render(scale=scale).to_pil().convert("RGB")
        buf = io.BytesIO()
        img.save(buf, "PNG")
        out.append((i + 1, buf.getvalue()))
    return out
