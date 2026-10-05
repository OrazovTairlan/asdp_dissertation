import io
import json

import docx
import pytest
from conftest import REPORT, SAMPLES, cond, make_catalog_xlsx
from openpyxl import load_workbook

from app import cli, docparse, pipeline
from app.catalog import build_catalog
from app.export import build_xlsx


def test_parse_txt_and_unsupported(tmp_path):
    p = tmp_path / "r.txt"
    p.write_text("Привет\nмир", encoding="utf-8")
    d = docparse.parse_file(p)
    assert d.text == "Привет\nмир" and d.tables == []
    with pytest.raises(ValueError, match="Неподдерживаемый"):
        docparse.parse_file(tmp_path / "x.rar")


def test_parse_xlsx_builds_catalog(tmp_path):
    d = docparse.parse_file(make_catalog_xlsx(tmp_path / "k.xlsx"))
    assert d.tables and d.tables[0].rows[0][0] == "Должность"
    cat = build_catalog([d])
    assert cat["warnings"] == [] and cat["positions"][0]["name"] == "Профессор"
    assert [len(b["items"]) for b in cat["positions"][0]["blocks"]] == [2, 2]


def test_parse_xlsx_merged_cells(tmp_path):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Должность", "Вид деятельности", "Вес", "Целевое значение"])
    ws.append(["1. Доцент", "Блок 1: Наука\nОдин из следующих:\n1. Статья\n2. Книга", "100", "1"])
    ws.merge_cells("A2:A3")
    ws.append([None, "Блок 2: Прочее\nОдин из следующих:\n1. Награда", "0", "1"])
    wb.save(tmp_path / "m.xlsx")
    d = docparse.parse_file(tmp_path / "m.xlsx")
    assert d.tables[0].rows[1][0].startswith("1. Доцент")


def test_parse_docx_with_table(tmp_path):
    doc = docx.Document()
    doc.add_paragraph("Для работников, не завершивших период")
    t = doc.add_table(rows=2, cols=4)
    for i, v in enumerate(["Должность", "Вид деятельности", "Вес", "Целевое значение"]):
        t.rows[0].cells[i].text = v
    for i, v in enumerate(["1. Доцент", "Блок 1: Наука\nОдин из следующих:\n1. Статья", "100", "1"]):
        t.rows[1].cells[i].text = v
    doc.save(tmp_path / "k.docx")
    d = docparse.parse_file(tmp_path / "k.docx")
    assert d.tables[0].context.startswith("Для работников")
    cat = build_catalog([d])
    assert cat["positions"][0]["name"] == "Доцент" and cat["positions"][0]["weight_sum"] == 100


def test_markdown_rendering_escapes_pipes_and_newlines():
    t = docparse.TableData(rows=[["a|b", "c\nd"], ["1", "2"]], source="s", page=1)
    md = t.to_markdown()
    assert "a\\|b" in md and "c<br>d" in md


def test_render_pdf_page_returns_png():
    pages = docparse.render_pdf_pages(SAMPLES / "kpi_polozhenie.pdf", [1], scale=0.5)
    assert pages[0][0] == 1 and pages[0][1][:4] == b"\x89PNG"


def test_catalog_without_tables_warns(tmp_path):
    p = tmp_path / "t.txt"
    p.write_text("нет таблиц", encoding="utf-8")
    cat = build_catalog([docparse.parse_file(p)])
    assert cat["positions"] == [] and any("Не удалось распознать" in w for w in cat["warnings"])


def test_weight_sum_warning(tmp_path):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Должность", "Вид деятельности", "Вес", "Целевое значение"])
    ws.append(["1. Доцент", "Блок 1: Наука\nОдин из следующих:\n1. Статья", "70", "1"])
    wb.save(tmp_path / "w.xlsx")
    cat = build_catalog([docparse.parse_file(tmp_path / "w.xlsx")])
    assert any("сумма весов" in w for w in cat["warnings"])


def _state():
    return {
        "result": {
            "employee": "Тестов Т.Т.",
            "position": {"name": "Профессор"},
            "weights_total": 100,
            "kpi_total_pct": 90.0,
            "blocks": [
                {
                    "title": "Научная работа",
                    "kind": "count",
                    "target": {"value": 1.0},
                    "achieved": 1,
                    "fulfillment": 100.0,
                    "weight": 60.0,
                    "contribution": 60.0,
                    "matches": [
                        {
                            "indicator_num": "1",
                            "quote": "Статья Alpha",
                            "source": "документ",
                            "counted": True,
                            "missing_conditions": "",
                        },
                        {
                            "indicator_num": "2",
                            "quote": "Доклад",
                            "source": "документ",
                            "counted": False,
                            "missing_conditions": "нужно 2",
                        },
                    ],
                    "metrics": [],
                    "notes": [],
                },
                {
                    "title": "Качество",
                    "kind": "metric",
                    "target": {"value": 80.0},
                    "achieved": 60.0,
                    "fulfillment": 75.0,
                    "weight": 40.0,
                    "contribution": 30.0,
                    "matches": [],
                    "metrics": [{"indicator_num": "1", "quote": "60%", "source": "документ"}],
                    "notes": ["заметка"],
                },
            ],
        }
    }


def test_export_layout():
    ws = load_workbook(io.BytesIO(build_xlsx(_state()))).active
    assert ws.max_row == 4
    assert [c.value for c in ws[1]][3] == "Наименование KPI"
    assert ws["E2"].value == "шт." and ws["E3"].value == "%"
    assert "НЕ ЗАСЧИТАНО [2]: нужно 2" in ws["L2"].value and "Статья Alpha" in ws["K2"].value
    assert ws["J4"].value == 90.0 and "Комиссией" in ws["L4"].value


def test_cli_prompts(capsys):
    assert cli.main(["prompts"]) == 0
    out = capsys.readouterr().out
    assert "PROMPT_VERSION" in out and "===== classify =====" in out


def test_cli_requires_input():
    with pytest.raises(SystemExit):
        cli.main(["evaluate", "--catalog", "x.xlsx"])


def test_cli_evaluate_end_to_end(tmp_path, fake_llm, capsys):
    cat = make_catalog_xlsx(tmp_path / "kpi.xlsx")
    rep = tmp_path / "report.txt"
    rep.write_text(REPORT, encoding="utf-8")
    fake_llm.achievements = [
        {"quote": "Статья «Alpha» принята в журнал Q1 по Scopus, аффилиация AITU, статус Accepted.", "summary": "с"}
    ]
    fake_llm.decisions = {
        "Alpha": {
            "indicator_id": "P1-B1-I1",
            "required_count": 1,
            "conditions": [cond("Q1", True, "журнал Q1"), cond("Accepted", True, "статус Accepted")],
            "reasoning": "",
        }
    }
    fake_llm.metrics = [{"indicator_id": "P1-B2-I1", "value": 60, "quote": "Результат анкетирования студентов: 60%."}]
    out_json = tmp_path / "out.json"
    code = cli.main(["evaluate", "--catalog", str(cat), "--report", str(rep), "--json", str(out_json)])
    assert code == 0
    printed = capsys.readouterr().out
    assert "ИТОГО: 90.0%" in printed
    assert json.loads(out_json.read_text(encoding="utf-8"))["result"]["kpi_total_pct"] == 90.0


def test_cli_reports_error(tmp_path, monkeypatch, capsys):
    cat = make_catalog_xlsx(tmp_path / "kpi.xlsx")
    rep = tmp_path / "r.txt"
    rep.write_text("текст", encoding="utf-8")
    monkeypatch.setattr(pipeline, "process", lambda sid, catalog, reuse=False: _fail(sid))
    assert cli.main(["evaluate", "--catalog", str(cat), "--report", str(rep)]) == 1


def _fail(sid):
    st = pipeline.load_state(sid)
    st.update(status="error", error="RuntimeError: boom")
    pipeline.save_state(st)
